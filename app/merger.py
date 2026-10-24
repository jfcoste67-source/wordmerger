import io
from copy import deepcopy
from pathlib import Path

from docxtpl import DocxTemplate
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TEMPLATE_DIR = Path(__file__).parent.parent / "templates"
ARTICLE2_DATE_FIELDS = {"DATE_DEBUT_REMPLACEMENT", "DATE_FIN_REMPLACEMENT"}
FR_MONTH_ABBR = {
    1: "janv.",
    2: "févr.",
    3: "mars",
    4: "avr.",
    5: "mai",
    6: "juin",
    7: "juil.",
    8: "août",
    9: "sept.",
    10: "oct.",
    11: "nov.",
    12: "déc.",
}
FR_MONTH_FULL = (
    "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
)
SLOT_MARKER = "__WORDMERGER_ACCEPTED_SLOTS__"
FIELD_ALIASES = {
    "MEDECIN_REMPMACE_CABINET_ADRESSE": "MEDECIN_REMPLACE_CABINET_ADRESSE",
    "MEDECIN_REMPMACE_CABINET_CODE_POSTAL": "MEDECIN_REMPLACE_CABINET_CODE_POSTAL",
    "MEDECIN_REMPMACE_CABINET_VILLE": "MEDECIN_REMPLACE_CABINET_VILLE",
    "MEDECIN_REMPLACANT_RRPS": "MEDECIN_REMPLACANT_RPPS",
    "MEDECIN_REMPLACE_RRPS": "MEDECIN_REMPLACE_RPPS",
}


def normalize_fields_with_aliases(fields: dict) -> dict[str, object]:
    normalized = {k.upper(): v for k, v in fields.items()}
    for alias, canonical in FIELD_ALIASES.items():
        if alias in normalized and canonical not in normalized:
            normalized[canonical] = normalized[alias]
        normalized.pop(alias, None)
    return normalized


def _replacement_defaults(fields: dict[str, object]) -> None:
    if "TAUX_CABINET" not in fields and "TAUX_RETROCESSION" in fields:
        try:
            rate = Decimal(str(fields["TAUX_RETROCESSION"]).strip().replace(",", "."))
        except InvalidOperation as exc:
            raise ValueError("TAUX_RETROCESSION must be a number between 0 and 100") from exc
        if not rate.is_finite() or not 0 <= rate <= 100:
            raise ValueError("TAUX_RETROCESSION must be a number between 0 and 100")
        fields["TAUX_CABINET"] = format(Decimal(100) - rate, "f")

    if "CRENEAUX_LISTE" not in fields and "JOURS_REMPLACEMENT" in fields:
        fields["CRENEAUX_LISTE"] = fields["JOURS_REMPLACEMENT"]

    if "AVERTISSEMENT_PROFIL" not in fields:
        status = str(fields.get("MEDECIN_REMPLACANT_THESE", "")).strip().casefold()
        if status in {"oui", "true", "1", "yes"}:
            fields["AVERTISSEMENT_PROFIL"] = ""
        elif status in {"non", "false", "0", "no"}:
            fields["AVERTISSEMENT_PROFIL"] = (
                "Profil non th\u00e9s\u00e9 : v\u00e9rifier que le contrat est adapt\u00e9 "
                "au statut du rempla\u00e7ant et que son autorisation de remplacement est valide."
            )
        else:
            fields["AVERTISSEMENT_PROFIL"] = (
                "Statut th\u00e9s\u00e9 du rempla\u00e7ant non renseign\u00e9 ou non reconnu : "
                "\u00e0 v\u00e9rifier avant signature."
            )
    if "DATE_JOUR" not in fields:
        current_date = datetime.now(ZoneInfo("Europe/Paris")).date()
        fields["DATE_JOUR"] = (
            f"{current_date.day} {FR_MONTH_FULL[current_date.month - 1]} {current_date.year}"
        )


def _format_date_fr_jj_mmm_aaaa(value: object) -> object:
    if isinstance(value, datetime):
        dt = value.date()
    elif isinstance(value, date):
        dt = value
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            return value
        dt = None
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                dt = datetime.strptime(text, fmt).date()
                break
            except ValueError:
                continue
        if dt is None:
            return value
    else:
        return value

    return f"{dt.day} {FR_MONTH_ABBR[dt.month]} {dt.year:04d}"


def _slot_lines(fields: dict[str, object]) -> list[str]:
    value = fields.get("CRENEAUX_ACCEPTES")
    if value is None:
        value = fields.get("CRENEAUX_LISTE", fields.get("JOURS_REMPLACEMENT", []))
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str):
        return [
            line.strip().removeprefix("•").strip()
            for line in value.splitlines()
            if line.strip()
        ]
    return [str(value).strip()] if value is not None and str(value).strip() else []


def _paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _paragraphs(cell)


def _replace_slot_marker(docx_bytes: bytes, slots: list[str]) -> bytes:
    document = Document(io.BytesIO(docx_bytes))
    matches = [paragraph for paragraph in _paragraphs(document) if paragraph.text == SLOT_MARKER]
    if len(matches) != 1:
        raise ValueError("The template must contain exactly one accepted-slots paragraph.")

    placeholder = matches[0]
    anchor = placeholder._p
    template_run = placeholder.runs[0]._r if placeholder.runs else None
    template_properties = template_run.find(qn("w:rPr")) if template_run is not None else None
    for slot in slots:
        paragraph = deepcopy(placeholder._p)
        for child in list(paragraph):
            if child.tag != qn("w:pPr"):
                paragraph.remove(child)
        run = OxmlElement("w:r")
        if template_properties is not None:
            run.append(deepcopy(template_properties))
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = f"• {slot}"
        run.append(text)
        paragraph.append(run)
        anchor.addnext(paragraph)
        anchor = paragraph
    placeholder._p.getparent().remove(placeholder._p)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def get_available_templates() -> list[str]:
    return [f.stem for f in sorted(TEMPLATE_DIR.glob("*.docx"))]


def merge(template_name: str, fields: dict) -> bytes:
    """Merge fields into a .docx template and return the document as bytes.

    Field keys are normalised to UPPERCASE to match {{FIELD}} placeholders.
    """
    template_path = TEMPLATE_DIR / f"{template_name}.docx"
    if not template_path.exists():
        raise FileNotFoundError(f"Template '{template_name}' not found")

    tpl = DocxTemplate(str(template_path))
    # Template placeholders use UPPERCASE — normalise incoming keys
    normalized = normalize_fields_with_aliases(fields)
    slots: list[str] = []
    if template_name == "contrat_remplacement":
        _replacement_defaults(normalized)
        slots = _slot_lines(normalized)
        normalized["CRENEAUX_ACCEPTES"] = SLOT_MARKER
    for key in ARTICLE2_DATE_FIELDS:
        if key in normalized:
            normalized[key] = _format_date_fr_jj_mmm_aaaa(normalized[key])
    # Convert list fields to comma-separated strings for plain {{}} placeholders
    for key, val in normalized.items():
        if isinstance(val, list):
            normalized[key] = ", ".join(str(item) for item in val)
    tpl.render(normalized, autoescape=False)

    buf = io.BytesIO()
    tpl.save(buf)
    result = buf.getvalue()
    if template_name == "contrat_remplacement":
        result = _replace_slot_marker(result, slots)
    return result
