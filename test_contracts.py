import io
import os
import unittest
from datetime import datetime
from unittest.mock import patch
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from docx import Document
from docxtpl import DocxTemplate
from fastapi import HTTPException
from lxml import etree

with patch.dict(os.environ, {"WORDMERGER_API_KEY": "local-test-key"}):
    from app.merger import (
        FR_MONTH_FULL,
        TEMPLATE_DIR,
        _replacement_defaults,
        merge,
        normalize_fields_with_aliases,
    )
    from app.models import MergeRequest
    from app.schema import get_allowed_fields, get_required_fields
    from app.main import merge_contract, merge_contract_pdf


def sample_fields():
    return {
        "medecin_remplacant_prenom": "Jeanne",
        "medecin_remplacant_nom": "Martin",
        "medecin_remplacant_rpps": "12345678901",
        "medecin_remplace_prenom": "Paul",
        "medecin_remplace_nom": "Durand",
        "medecin_remplace_rpps": "10987654321",
        "medecin_remplace_cabinet_nom": "Cabinet Test",
        "medecin_remplace_cabinet_adresse": "12 rue des Fleurs",
        "medecin_remplace_cabinet_code_postal": "75015",
        "medecin_remplace_cabinet_ville": "Paris",
        "jours_remplacement": "1,5 jours",
        "creneaux_acceptes": ["21 déc. 2026 journée", "25 déc. 2026 matin"],
        "date_debut_remplacement": "2026-04-20",
        "date_fin_remplacement": "2026-06-30",
        "taux_retrocession": "70",
    }


class ContractTests(unittest.TestCase):
    def render_document(self, fields):
        response = merge_contract(
            MergeRequest(template="contrat_remplacement", fields=fields)
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment;", response.headers["content-disposition"])
        data = response.body
        with ZipFile(io.BytesIO(data)) as package:
            for name in package.namelist():
                if name.endswith(".xml"):
                    etree.fromstring(package.read(name))
        document = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertNotIn("{{", text)
        self.assertNotIn("{%", text)
        return document

    def render(self, fields):
        return "\n".join(p.text for p in self.render_document(fields).paragraphs)

    def test_template_schema_alignment(self):
        template = DocxTemplate(str(TEMPLATE_DIR / "contrat_remplacement.docx"))
        variables = template.get_undeclared_template_variables()
        self.assertLessEqual(variables, get_allowed_fields("contrat_remplacement"))
        self.assertLessEqual(get_required_fields("contrat_remplacement"), variables)

    def test_current_template_dates_slots_and_defaults(self):
        fields = sample_fields()
        original = fields.copy()
        text = self.render(fields)
        self.assertIn("12345678901", text)
        self.assertIn("10987654321", text)
        self.assertIn("20 avr. 2026", text)
        self.assertIn("30 juin 2026", text)
        today = datetime.now(ZoneInfo("Europe/Paris")).date()
        self.assertIn(f"{today.day} {FR_MONTH_FULL[today.month - 1]} {today.year}", text)
        self.assertIn("1,5 jours", text)
        self.assertIn("70 %", text)
        self.assertNotIn("30 %", text)
        self.assertIn("• 21 déc. 2026 journée", text)
        self.assertIn("• 25 déc. 2026 matin", text)
        self.assertIn("non renseign\u00e9 ou non reconnu", text)
        self.assertEqual(fields, original)

    def test_current_template_preserves_requested_formatting(self):
        document = self.render_document(sample_fields())
        paragraphs = document.paragraphs
        summary = next(p for p in paragraphs if p.text.startswith("Le présent contrat de remplacement"))
        self.assertEqual(
            "".join(run.text for run in summary.runs if run.bold),
            "pour une période de 1,5 jours s'étendant du 20 avr. 2026 au 30 juin 2026 compris",
        )
        self.assertEqual(
            [p.text for p in paragraphs if p.text.startswith("• ")],
            ["• 21 déc. 2026 journée", "• 25 déc. 2026 matin"],
        )

        article7 = next(p for p in paragraphs if "La part des honoraires revenant au remplaçant" in p.text)
        self.assertIn("La part des honoraires revenant au remplaçant est de 70 %.", article7.text)
        self.assertNotIn("du total des honoraires perçus", article7.text)
        self.assertTrue(any(run.bold and "70 %" in run.text for run in article7.runs))

        article9 = next(p for p in paragraphs if p.text.startswith("Conciliation :"))
        underlined = [run.text for run in article9.runs if run.underline]
        self.assertEqual("".join(underlined), "Conciliation :")

        notice = next(p for p in paragraphs if p.text.startswith("La plateforme ne recueille"))
        self.assertTrue(notice.runs[0].italic)
        self.assertEqual(notice.runs[0].font.size.pt, 10)
        signature_text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
        self.assertIn("Signature du médecin remplacé", signature_text)
        self.assertIn("Signature du remplaçant", signature_text)

    def test_complete_legacy_payload(self):
        fields = sample_fields()
        for role in ("remplacant", "remplace"):
            fields[f"medecin_{role}_rrps"] = fields.pop(f"medecin_{role}_rpps")
        fields.update({
            "medecin_remplacant_sexe": "F",
            "medecin_remplacant_these": "Oui",
            "medecin_remplacant_adresse": "8 avenue Victor Hugo",
            "medecin_remplacant_code_postal": "75016",
            "medecin_remplacant_ville": "Paris",
            "annonce_id": "A-2026-1001",
            "candidature_id": "C-2026-778",
            "cabinet_medecin_id": "CM-55",
            "statut_contrat_remplacement": "Signe",
            "date_signature_cabinet": "14/04/2026",
            "ip_signature": "127.0.0.1",
        })
        fields["medecin_rempmace_cabinet_adresse"] = fields.pop(
            "medecin_remplace_cabinet_adresse"
        )
        text = self.render(fields)
        self.assertIn("12345678901", text)
        self.assertIn("10987654321", text)
        self.assertIn("12 rue des Fleurs", text)
        self.assertNotIn("Profil non", text)
        self.assertNotIn("non renseign\u00e9", text)

    def test_direct_merge_supports_legacy_aliases(self):
        fields = sample_fields()
        fields["medecin_remplacant_rrps"] = fields.pop("medecin_remplacant_rpps")
        document = Document(io.BytesIO(merge("contrat_remplacement", fields)))
        self.assertIn("12345678901", "\n".join(p.text for p in document.paragraphs))

    def test_canonical_alias_precedence(self):
        normalized = normalize_fields_with_aliases({
            "medecin_remplacant_rrps": "legacy",
            "MEDECIN_REMPLACANT_RPPS": "canonical",
            "medecin_rempmace_cabinet_ville": "Paris",
            "medecin_rempmace_cabinet_code_postal": "75015",
        })
        self.assertEqual(normalized["MEDECIN_REMPLACANT_RPPS"], "canonical")
        self.assertNotIn("MEDECIN_REMPLACANT_RRPS", normalized)
        self.assertEqual(normalized["MEDECIN_REMPLACE_CABINET_VILLE"], "Paris")
        self.assertEqual(normalized["MEDECIN_REMPLACE_CABINET_CODE_POSTAL"], "75015")

    def test_explicit_values_override_defaults(self):
        fields = sample_fields()
        fields.update({
            "taux_cabinet": "25",
            "creneaux_liste": ["Lundi matin", "Mardi soir"],
            "creneaux_acceptes": ["Lundi matin", "Mardi soir"],
            "avertissement_profil": "Verification manuelle.",
            "medecin_remplacant_these": "non",
        })
        text = self.render(fields)
        self.assertNotIn("30 %", text)
        self.assertIn("• Lundi matin", text)
        self.assertIn("• Mardi soir", text)
        self.assertIn("Verification manuelle.", text)
        self.assertNotIn("Profil non", text)
        defaults = normalize_fields_with_aliases(fields)
        _replacement_defaults(defaults)
        self.assertEqual(defaults["TAUX_CABINET"], "25")

    def test_text_slots_and_explicit_empty_warning(self):
        fields = sample_fields()
        fields.update(creneaux_acceptes="Lundi matin", avertissement_profil="")
        text = self.render(fields)
        self.assertIn("Lundi matin", text)
        self.assertNotIn("non renseign\u00e9", text)

    def test_pdf_endpoint_returns_converted_pdf(self):
        with patch("app.main._convert_docx_to_pdf", return_value=b"%PDF-1.7\ncontract"):
            response = merge_contract_pdf(
                MergeRequest(template="contrat_remplacement", fields=sample_fields())
            )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.media_type, "application/pdf")
        self.assertIn("contrat_contrat_remplacement.pdf", response.headers["content-disposition"])
        self.assertTrue(response.body.startswith(b"%PDF-"))

    def test_profile_statuses(self):
        for status in ("non", " NON ", False, "false", 0):
            with self.subTest(status=status):
                fields = sample_fields()
                fields["medecin_remplacant_these"] = status
                self.assertIn("Profil non th\u00e9s\u00e9", self.render(fields))
        for status in ("oui", " OUI ", True, "true", 1):
            with self.subTest(status=status):
                fields = sample_fields()
                fields["medecin_remplacant_these"] = status
                text = self.render(fields)
                self.assertNotIn("Profil non", text)
                self.assertNotIn("non renseign\u00e9", text)
        fields = sample_fields()
        fields["medecin_remplacant_these"] = "unknown"
        self.assertIn("non renseign\u00e9 ou non reconnu", self.render(fields))

    def test_decimal_and_boundary_rates(self):
        for rate, expected in (("70.5", "29.5"), ("70,5", "29.5"), (0, "100"), (100, "0")):
            with self.subTest(rate=rate):
                fields = {"TAUX_RETROCESSION": rate}
                _replacement_defaults(fields)
                self.assertEqual(fields["TAUX_CABINET"], expected)

    def test_invalid_derived_rates_return_422(self):
        for rate in ("bad", "", None, True, "-1", "101", "NaN", "Infinity"):
            with self.subTest(rate=rate):
                fields = sample_fields()
                fields["taux_retrocession"] = rate
                with self.assertRaises(HTTPException) as caught:
                    self.render(fields)
                self.assertEqual(caught.exception.status_code, 422)
                self.assertIn("TAUX_RETROCESSION", caught.exception.detail)

    def test_missing_required_field_returns_422(self):
        fields = sample_fields()
        fields.pop("medecin_remplacant_rpps")
        with self.assertRaises(HTTPException) as caught:
            self.render(fields)
        self.assertEqual(caught.exception.status_code, 422)
        self.assertIn("MEDECIN_REMPLACANT_RPPS", caught.exception.detail["fields"])

    def test_unknown_field_still_returns_422(self):
        fields = sample_fields()
        fields["unknown_field"] = "unexpected"
        with self.assertRaises(HTTPException) as caught:
            self.render(fields)
        self.assertEqual(caught.exception.status_code, 422)
        self.assertEqual(caught.exception.detail["fields"], ["UNKNOWN_FIELD"])


if __name__ == "__main__":
    unittest.main()
